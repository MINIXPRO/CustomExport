import io
import frappe


# ---------------------------------------------------------------------------
# Template fields for Sales Order Sub Item (custom_sub_items child table).
# Doctype: "Sales Order Sub Item" (singular) — 20 fields in DB.
# Template includes the 10 business-relevant fields per user requirement.
# Excluded: parent_item_name, sub_item_name (auto-populated read-only),
#           base_rate, CIF fields (calculated), parent_row_uid (internal).
# ---------------------------------------------------------------------------
TEMPLATE_FIELDS = [
    ("Parent Item",                    "parent_item"),
    ("Sub Item Code",                  "sub_item_code"),
    ("Sub Description",                "sub_description"),
    ("HSN/SAC",                        "gst_hsn_code"),
    ("Delivery Date",                  "delivery_date"),
    ("Quantity",                       "qty"),
    ("Rate",                           "rate"),
    ("UOM",                            "uom"),
    ("Distributed Discount Amount",    "custom_distributed_discount_amount"),
    ("Amount",                         "amount"),
]


@frappe.whitelist()
def get_so_sub_items_template(sales_order=None):
    """
    Generate and return an XLSX template for Sales Order Sub Items
    bulk import/export.

    Row 1 → column labels  (human-readable, styled dark-blue)
    Row 2 → fieldnames     (machine-readable, styled mid-blue)
    Row 3+ → existing sub item data when `sales_order` is provided.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        frappe.throw("openpyxl is required. Run: pip install openpyxl")

    wb = Workbook()
    ws = wb.active
    ws.title = "SO Sub Items"

    header_font    = Font(bold=True, color="FFFFFF")
    header_fill    = PatternFill(start_color="2E4057", end_color="2E4057", fill_type="solid")
    fieldname_fill = PatternFill(start_color="5B8DB8", end_color="5B8DB8", fill_type="solid")
    fieldname_font = Font(bold=False, color="FFFFFF")
    center_align   = Alignment(horizontal="center", vertical="center", wrap_text=True)

    labels     = [label     for label, _  in TEMPLATE_FIELDS]
    fieldnames = [fieldname for _, fieldname in TEMPLATE_FIELDS]

    ws.append(labels)
    ws.append(fieldnames)

    for col_idx, _ in enumerate(TEMPLATE_FIELDS, start=1):
        cell_label = ws.cell(row=1, column=col_idx)
        cell_label.font      = header_font
        cell_label.fill      = header_fill
        cell_label.alignment = center_align

        cell_field = ws.cell(row=2, column=col_idx)
        cell_field.font      = fieldname_font
        cell_field.fill      = fieldname_fill
        cell_field.alignment = center_align

        max_len = max(len(labels[col_idx - 1]), len(fieldnames[col_idx - 1]))
        ws.column_dimensions[
            ws.cell(row=1, column=col_idx).column_letter
        ].width = max(14, min(max_len + 2, 40))

    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 22

    if sales_order:
        try:
            so = frappe.get_doc("Sales Order", sales_order)
            for sub in (so.custom_sub_items or []):
                row = []
                for _, fn in TEMPLATE_FIELDS:
                    val = getattr(sub, fn, None)
                    row.append("" if val is None else val)
                ws.append(row)
        except frappe.DoesNotExistError:
            pass

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    fname = (
        f"so_sub_items_{sales_order}.xlsx"
        if sales_order
        else "so_sub_items_template.xlsx"
    )
    frappe.response["filename"]    = fname
    frappe.response["filecontent"] = output.read()
    frappe.response["type"]        = "download"


@frappe.whitelist()
def import_so_sub_items(sales_order, file_url):
    """
    Import Sales Order Sub Item rows from the custom XLSX / CSV template.

    STRATEGY: Clear and Rebuild.
    Uploaded rows are the single source of truth for custom_sub_items.

    Row 1: labels      (ignored)
    Row 2: fieldnames  (column-to-field mapping)
    Row 3+: data rows  → become the complete new custom_sub_items

    Each row must have parent_item + sub_item_code as identifiers.
    parent_row_uid is resolved automatically from the parent item row.
    """
    import os
    from collections import defaultdict

    site_path = frappe.get_site_path()
    if file_url.startswith("/files/"):
        file_path = os.path.join(site_path, "public", file_url.lstrip("/"))
    elif file_url.startswith("/private/files/"):
        file_path = os.path.join(site_path, file_url.lstrip("/"))
    else:
        frappe.throw(f"Unsupported file URL format: {file_url}")

    if not os.path.exists(file_path):
        frappe.throw(f"File not found on disk: {file_path}")

    ext = os.path.splitext(file_url.lower())[1]
    if ext in (".xlsx", ".xls"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            frappe.throw("openpyxl is required. Run: pip install openpyxl")
        wb   = load_workbook(file_path, data_only=True)
        rows = list(wb.active.iter_rows(values_only=True))
    elif ext == ".csv":
        import csv
        with open(file_path, newline="", encoding="utf-8-sig") as f:
            rows = [tuple(r) for r in csv.reader(f)]
    else:
        frappe.throw(f"Unsupported file format '{ext}'. Upload a .xlsx or .csv file.")

    if len(rows) < 3:
        frappe.throw(
            "The uploaded file must have at least 3 rows "
            "(Row 1 = labels, Row 2 = fieldnames, Row 3+ = data)."
        )

    fieldnames = [str(f).strip() if f is not None else None for f in rows[1]]

    SKIP_FIELDS = {
        "sub_item_code", "parent_item", "parent_row_uid",
        "name", "parent", "parenttype", "parentfield", "idx",
        "parent_item_name", "sub_item_name",
    }

    parsed_rows = []
    for data_row in rows[2:]:
        if all(v is None or str(v).strip() == "" for v in data_row):
            continue
        row_data = {}
        for col_idx, fn in enumerate(fieldnames):
            if not fn or fn == "None":
                continue
            row_data[fn] = data_row[col_idx] if col_idx < len(data_row) else None

        sub_item_code = str(row_data.get("sub_item_code") or "").strip()
        parent_item   = str(row_data.get("parent_item") or "").strip()
        if not sub_item_code or not parent_item:
            continue
        parsed_rows.append((parent_item, sub_item_code, row_data))

    if not parsed_rows:
        frappe.throw("No valid sub item rows found. Each row needs parent_item and sub_item_code.")

    so = frappe.get_doc("Sales Order", sales_order)
    frappe.has_permission("Sales Order", "write", so, throw=True)

    # Build parent item → parent_row_uid map
    parent_uid_map = {}
    for item in (so.items or []):
        code = str(item.item_code or "").strip()
        if code and code not in parent_uid_map:
            uid = getattr(item, "custom_row_uid", None) or item.name
            parent_uid_map[code] = uid

    # Build existing sub-items lookup by (parent_item, sub_item_code)
    existing_map = defaultdict(list)
    for sub in (so.custom_sub_items or []):
        key = (
            str(sub.parent_item   or "").strip(),
            str(sub.sub_item_code or "").strip(),
        )
        existing_map[key].append(sub)

    so.custom_sub_items = []
    usage_counter = defaultdict(int)

    for parent_item, sub_item_code, row_data in parsed_rows:
        key = (parent_item, sub_item_code)
        candidates = existing_map.get(key, [])
        idx = usage_counter[key]
        usage_counter[key] += 1

        if idx < len(candidates):
            base = candidates[idx].as_dict()
            base.pop("name", None)
            base.pop("idx", None)
            new_sub = so.append("custom_sub_items", base)
        else:
            new_sub = so.append("custom_sub_items", {
                "parent_item":   parent_item,
                "sub_item_code": sub_item_code,
            })

        new_sub.parent_row_uid = parent_uid_map.get(parent_item, "")

        for fn, value in row_data.items():
            if fn in SKIP_FIELDS:
                continue
            if not hasattr(new_sub, fn):
                continue
            setattr(new_sub, fn, None if (value is None or str(value).strip() == "") else value)

    total = len(parsed_rows)
    so.save(ignore_permissions=True)
    frappe.db.commit()

    return {
        "total":   total,
        "message": f"Successfully imported {total} sub item row(s) into {sales_order}. Sub items replaced.",
    }
