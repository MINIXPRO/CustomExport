import io
import frappe


# ---------------------------------------------------------------------------
# Template fields for Packing Slip Item bulk import/export.
# Excluded: calculated CIF fields (recomputed on save), custom_row_uid (internal),
#           custom_currency (inherited), custom_rate (duplicate of rate).
# ---------------------------------------------------------------------------
TEMPLATE_FIELDS = [
    ("Item Code",              "item_code"),
    ("Item Name",              "item_name"),
    ("Sales Order No",         "custom_sales_order_no"),
    ("Cust PO No",             "custom_cust_po_no"),
    ("Customer Part No",       "custom_customer_part_no"),
    ("Quantity",               "qty"),
    ("Rate",                   "rate"),
    ("Unit Weight (Kg)",       "custom_unit_weight"),
    ("Gross Weight (Kg)",      "custom__gross_weight"),
    ("Box",                    "custom_box"),
    ("Box Type",               "custom_box_type"),
    ("Length (Inch)",          "custom_length"),
    ("Width (Inch)",           "custom_width"),
    ("Height (Inch)",          "custom_height"),
    ("Cubic Feet",             "custom_cubic_feet"),
    ("Cubic Meter",            "custom_cubic_meter"),
    ("Freight & Insurance %",  "custom_freight__insurance_"),
]


@frappe.whitelist()
def get_ps_items_template(packing_slip=None):
    """
    Generate and return an XLSX template for Packing Slip Item bulk import/export.

    Row 1 → column labels  (human-readable, styled dark-blue)
    Row 2 → fieldnames     (machine-readable, styled mid-blue)
    Row 3+ → existing item data when `packing_slip` is provided.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        frappe.throw("openpyxl is required. Run: pip install openpyxl")

    wb = Workbook()
    ws = wb.active
    ws.title = "PS Items"

    header_font    = Font(bold=True, color="FFFFFF")
    header_fill    = PatternFill(start_color="2E4057", end_color="2E4057", fill_type="solid")
    fieldname_fill = PatternFill(start_color="5B8DB8", end_color="5B8DB8", fill_type="solid")
    fieldname_font = Font(bold=False, color="FFFFFF")
    center_align   = Alignment(horizontal="center", vertical="center", wrap_text=True)

    labels     = [label     for label, _        in TEMPLATE_FIELDS]
    fieldnames = [fieldname for _, fieldname in TEMPLATE_FIELDS]

    ws.append(labels)
    ws.append(fieldnames)

    for col_idx, _ in enumerate(TEMPLATE_FIELDS, start=1):
        cell_label = ws.cell(row=1, column=col_idx)
        cell_label.font      = header_font
        cell_label.fill      = header_fill
        cell_label.alignment = center_align

        cell_field = ws.cell(row=2, column=col_idx)
        cell_field.font      = fieldname_font
        cell_field.fill      = fieldname_fill
        cell_field.alignment = center_align

        max_len = max(len(labels[col_idx - 1]), len(fieldnames[col_idx - 1]))
        ws.column_dimensions[
            ws.cell(row=1, column=col_idx).column_letter
        ].width = max(14, min(max_len + 2, 40))

    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 22

    if packing_slip:
        try:
            ps = frappe.get_doc("Packing Slip", packing_slip)
            for item in (ps.items or []):
                row = []
                for _, fn in TEMPLATE_FIELDS:
                    val = getattr(item, fn, None)
                    row.append("" if val is None else val)
                ws.append(row)
        except frappe.DoesNotExistError:
            pass

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    fname = (
        f"ps_items_{packing_slip}.xlsx"
        if packing_slip
        else "ps_items_template.xlsx"
    )
    frappe.response["filename"]    = fname
    frappe.response["filecontent"] = output.read()
    frappe.response["type"]        = "download"


@frappe.whitelist()
def import_ps_items(packing_slip, file_url):
    """
    Import Packing Slip Item rows from the custom XLSX / CSV template.

    STRATEGY: Clear and Rebuild.
    Uploaded rows are the single source of truth for ps.items.

    Row 1: labels      (ignored)
    Row 2: fieldnames  (column-to-field mapping)
    Row 3+: data rows  → become the complete new ps.items

    Existing rows are matched by item_code to preserve linked fields
    (warehouse, SO reference, etc.). Template values are applied on top.
    """
    import os
    from collections import defaultdict

    site_path = frappe.get_site_path()
    if file_url.startswith("/files/"):
        file_path = os.path.join(site_path, "public", file_url.lstrip("/"))
    elif file_url.startswith("/private/files/"):
        file_path = os.path.join(site_path, file_url.lstrip("/"))
    else:
        frappe.throw(f"Unsupported file URL format: {file_url}")

    if not os.path.exists(file_path):
        frappe.throw(f"File not found on disk: {file_path}")

    ext = os.path.splitext(file_url.lower())[1]
    if ext in (".xlsx", ".xls"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            frappe.throw("openpyxl is required. Run: pip install openpyxl")
        wb   = load_workbook(file_path, data_only=True)
        rows = list(wb.active.iter_rows(values_only=True))
    elif ext == ".csv":
        import csv
        with open(file_path, newline="", encoding="utf-8-sig") as f:
            rows = [tuple(r) for r in csv.reader(f)]
    else:
        frappe.throw(f"Unsupported file format '{ext}'. Upload a .xlsx or .csv file.")

    if len(rows) < 3:
        frappe.throw(
            "The uploaded file must have at least 3 rows "
            "(Row 1 = labels, Row 2 = fieldnames, Row 3+ = data)."
        )

    fieldnames = [str(f).strip() if f is not None else None for f in rows[1]]

    SKIP_FIELDS = {"item_code", "name", "parent", "parenttype", "parentfield", "idx"}

    parsed_rows = []
    for data_row in rows[2:]:
        if all(v is None or str(v).strip() == "" for v in data_row):
            continue
        row_data = {}
        for col_idx, fn in enumerate(fieldnames):
            if not fn or fn == "None":
                continue
            row_data[fn] = data_row[col_idx] if col_idx < len(data_row) else None
        item_code = str(row_data.get("item_code") or "").strip()
        if not item_code:
            continue
        parsed_rows.append((item_code, row_data))

    if not parsed_rows:
        frappe.throw("No valid item rows found in the uploaded file.")

    ps = frappe.get_doc("Packing Slip", packing_slip)
    frappe.has_permission("Packing Slip", "write", ps, throw=True)

    existing_map = defaultdict(list)
    for item in (ps.items or []):
        existing_map[str(item.item_code).strip()].append(item)

    ps.items = []
    usage_counter = defaultdict(int)

    for item_code, row_data in parsed_rows:
        candidates = existing_map.get(item_code, [])
        idx = usage_counter[item_code]
        usage_counter[item_code] += 1

        if idx < len(candidates):
            base = candidates[idx].as_dict()
            base.pop("name", None)
            base.pop("idx", None)
            new_item = ps.append("items", base)
        else:
            new_item = ps.append("items", {"item_code": item_code})

        for fn, value in row_data.items():
            if fn in SKIP_FIELDS:
                continue
            if not hasattr(new_item, fn):
                continue
            setattr(new_item, fn, None if (value is None or str(value).strip() == "") else value)

    total = len(parsed_rows)
    ps.save(ignore_permissions=True)
    frappe.db.commit()

    return {
        "total":   total,
        "message": f"Successfully imported {total} item row(s) into {packing_slip}. Packing Slip items replaced.",
    }
