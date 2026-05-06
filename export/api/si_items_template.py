import io
import frappe


# ---------------------------------------------------------------------------
# Template fields for Sales Invoice Item bulk import/export.
# CIF calculated fields (custom_cif_unit_price*, custom___cif_total_amount*,
# custom__cif_total_amount*) are read-only on SI Item and excluded here —
# they are recalculated on save. Internal field custom_row_uid also excluded.
# ---------------------------------------------------------------------------
TEMPLATE_FIELDS = [
    ("Customer PO Number",     "custom_customer_order_number"),
    ("Sales Order",            "sales_order"),
    ("Item Code",              "item_code"),
    ("Item Name",              "item_name"),
    ("Customer Item Code",     "customer_item_code"),
    ("Description",            "description"),
    ("Qty",                    "qty"),
    ("UOM",                    "uom"),
    ("Rate",                   "rate"),
    ("Amount",                 "amount"),
    ("HSN/SAC",                "gst_hsn_code"),
    ("Freight & Insurance %",  "custom_freight__insurance_"),
    ("Duty Drawback",          "custom_duty_drawback"),
    ("Total Weight",           "total_weight"),
    ("Sales Order Item ID",    "so_detail"),
]


@frappe.whitelist()
def get_si_items_template(sales_invoice=None):
    """
    Generate and return an XLSX template for Sales Invoice Item bulk import/export.

    Row 1 → column labels  (human-readable, styled dark-blue)
    Row 2 → fieldnames     (machine-readable, styled mid-blue)
    Row 3+ → existing item data when `sales_invoice` is provided.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        frappe.throw("openpyxl is required. Run: pip install openpyxl")

    wb = Workbook()
    ws = wb.active
    ws.title = "SI Items"

    header_font    = Font(bold=True, color="FFFFFF")
    header_fill    = PatternFill(start_color="2E4057", end_color="2E4057", fill_type="solid")
    fieldname_fill = PatternFill(start_color="5B8DB8", end_color="5B8DB8", fill_type="solid")
    fieldname_font = Font(bold=False, color="FFFFFF")
    center_align   = Alignment(horizontal="center", vertical="center", wrap_text=True)

    labels     = [label     for label, _        in TEMPLATE_FIELDS]
    fieldnames = [fieldname for _, fieldname in TEMPLATE_FIELDS]

    ws.append(labels)
    ws.append(fieldnames)

    for col_idx, _ in enumerate(TEMPLATE_FIELDS, start=1):
        cell_label = ws.cell(row=1, column=col_idx)
        cell_label.font      = header_font
        cell_label.fill      = header_fill
        cell_label.alignment = center_align

        cell_field = ws.cell(row=2, column=col_idx)
        cell_field.font      = fieldname_font
        cell_field.fill      = fieldname_fill
        cell_field.alignment = center_align

        max_len = max(len(labels[col_idx - 1]), len(fieldnames[col_idx - 1]))
        ws.column_dimensions[
            ws.cell(row=1, column=col_idx).column_letter
        ].width = max(14, min(max_len + 2, 40))

    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 22

    if sales_invoice:
        try:
            si = frappe.get_doc("Sales Invoice", sales_invoice)
            for item in (si.items or []):
                row = []
                for _, fn in TEMPLATE_FIELDS:
                    val = getattr(item, fn, None)
                    row.append("" if val is None else val)
                ws.append(row)
        except frappe.DoesNotExistError:
            pass

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    fname = (
        f"si_items_{sales_invoice}.xlsx"
        if sales_invoice
        else "si_items_template.xlsx"
    )
    frappe.response["filename"]    = fname
    frappe.response["filecontent"] = output.read()
    frappe.response["type"]        = "download"


@frappe.whitelist()
def import_si_items(sales_invoice, file_url):
    import os
    from collections import defaultdict

    site_path = frappe.get_site_path()
    if file_url.startswith("/files/"):
        file_path = os.path.join(site_path, "public", file_url.lstrip("/"))
    elif file_url.startswith("/private/files/"):
        file_path = os.path.join(site_path, file_url.lstrip("/"))
    else:
        frappe.throw(f"Unsupported file URL format: {file_url}")

    if not os.path.exists(file_path):
        frappe.throw(f"File not found on disk: {file_path}")

    ext = os.path.splitext(file_url.lower())[1]
    if ext in (".xlsx", ".xls"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            frappe.throw("openpyxl is required. Run: pip install openpyxl")
        wb = load_workbook(file_path, data_only=True)
        rows = list(wb.active.iter_rows(values_only=True))
    elif ext == ".csv":
        import csv
        with open(file_path, newline="", encoding="utf-8-sig") as f:
            rows = [tuple(r) for r in csv.reader(f)]
    else:
        frappe.throw(f"Unsupported file format '{ext}'.")

    if len(rows) < 3:
        frappe.throw("File must have at least 3 rows.")

    fieldnames = [str(f).strip() if f is not None else None for f in rows[1]]

    SKIP_FIELDS = {"item_code", "name", "parent", "parenttype", "parentfield", "idx"}

    parsed_rows = []
    for data_row in rows[2:]:
        if all(v is None or str(v).strip() == "" for v in data_row):
            continue

        row_data = {}
        for col_idx, fn in enumerate(fieldnames):
            if not fn or fn == "None":
                continue
            row_data[fn] = data_row[col_idx] if col_idx < len(data_row) else None

        item_code = str(row_data.get("item_code") or "").strip()
        if not item_code:
            continue

        parsed_rows.append((item_code, row_data))

    if not parsed_rows:
        frappe.throw("No valid rows found in file.")

    si = frappe.get_doc("Sales Invoice", sales_invoice)
    frappe.has_permission("Sales Invoice", "write", si, throw=True)

    existing_map = defaultdict(list)
    for item in (si.items or []):
        existing_map[str(item.item_code).strip()].append(item)

    si.items = []
    usage_counter = defaultdict(int)

    # ✅ Collect issues
    invalid_rows = []              # PO mismatch
    missing_so_detail_rows = []   # Missing so_detail

    for item_code, row_data in parsed_rows:
        so_name = row_data.get("sales_order")
        customer_po = str(row_data.get("custom_customer_order_number") or "").strip()

        # ✅ FIRST: SO DETAIL VALIDATION
        so_detail = str(row_data.get("so_detail") or "").strip()
        if not so_detail:
            missing_so_detail_rows.append({
                "item_code": item_code,
                "sales_order": so_name or "",
                "row_po": customer_po
            })
            continue  # ❌ skip

        # ✅ SECOND: PO VALIDATION (unchanged)
        if so_name:
            try:
                so_doc = frappe.get_doc("Sales Order", so_name)
                so_po = str(getattr(so_doc, "po_no", "") or "").strip()

                if so_po and customer_po and so_po != customer_po:
                    invalid_rows.append({
                        "item_code": item_code,
                        "sales_order": so_name,
                        "so_po": so_po,
                        "row_po": customer_po
                    })
                    continue

            except frappe.DoesNotExistError:
                invalid_rows.append({
                    "item_code": item_code,
                    "sales_order": so_name,
                    "so_po": "NOT FOUND",
                    "row_po": customer_po
                })
                continue

        # ✅ EXISTING LOGIC (UNCHANGED)
        candidates = existing_map.get(item_code, [])
        idx = usage_counter[item_code]
        usage_counter[item_code] += 1

        if idx < len(candidates):
            base = candidates[idx].as_dict()
            base.pop("name", None)
            base.pop("idx", None)
            new_item = si.append("items", base)
        else:
            new_item = si.append("items", {"item_code": item_code})

        for fn, value in row_data.items():
            if fn in SKIP_FIELDS:
                continue
            if not hasattr(new_item, fn):
                continue
            setattr(
                new_item,
                fn,
                None if (value is None or str(value).strip() == "") else value
            )

    total = len(si.items)

    si.save(ignore_permissions=True)
    frappe.db.commit()

    # ✅ POPUP MESSAGE
    if invalid_rows or missing_so_detail_rows:
        message = '<div style="max-height:300px; overflow:auto;">'

        # 🔴 Missing SO Detail Table
        if missing_so_detail_rows:
            message += """
                <p style="margin:10px 0;"><b>Rows skipped due to missing Sales Order Item ID:</b></p>
                <table style="width:100%; border-collapse:separate; border-spacing:0; font-size:13px;">
                <thead>
                    <tr style="background-color:#fff3cd;">
                        <th style="padding:8px; border:1px solid #d1d8dd;">Item Code</th>
                        <th style="padding:8px; border:1px solid #d1d8dd;">Sales Order</th>
                        <th style="padding:8px; border:1px solid #d1d8dd;">File PO</th>
                    </tr>
                </thead>
                <tbody>
            """
            for r in missing_so_detail_rows:
                message += f"""
                    <tr>
                        <td style="padding:8px; border:1px solid #e4e7eb;">{r['item_code']}</td>
                        <td style="padding:8px; border:1px solid #e4e7eb;">{r['sales_order']}</td>
                        <td style="padding:8px; border:1px solid #e4e7eb;">{r['row_po']}</td>
                    </tr>
                """
            message += "</tbody></table>"

        # 🔵 PO Mismatch Table (your original)
        if invalid_rows:
            message += """
                <p style="margin:10px 0;"><b>Rows skipped due to PO mismatch:</b></p>
                <table style="width:100%; border-collapse:separate; border-spacing:0; font-size:13px;">
                <thead>
                    <tr style="background-color:#f7fafc;">
                        <th style="padding:8px; border:1px solid #d1d8dd;">Item Code</th>
                        <th style="padding:8px; border:1px solid #d1d8dd;">Sales Order</th>
                        <th style="padding:8px; border:1px solid #d1d8dd;">SO PO</th>
                        <th style="padding:8px; border:1px solid #d1d8dd;">File PO</th>
                    </tr>
                </thead>
                <tbody>
            """
            for r in invalid_rows:
                message += f"""
                    <tr>
                        <td style="padding:8px; border:1px solid #e4e7eb;">{r['item_code']}</td>
                        <td style="padding:8px; border:1px solid #e4e7eb;">{r['sales_order']}</td>
                        <td style="padding:8px; border:1px solid #e4e7eb; color:#d9534f;">{r['so_po']}</td>
                        <td style="padding:8px; border:1px solid #e4e7eb; color:#5bc0de;">{r['row_po']}</td>
                    </tr>
                """
            message += "</tbody></table>"

        message += "</div>"
        frappe.msgprint(message)

    return {
        "total": total,
        "message": f"Imported {total} valid row(s). "
                   f"{len(missing_so_detail_rows)} missing SO detail, "
                   f"{len(invalid_rows)} PO mismatch."
    }